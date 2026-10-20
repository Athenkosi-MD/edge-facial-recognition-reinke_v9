"""
app/auth.py  – Authentication, User-Management & Custom Fields Blueprint
"""

from flask import (Blueprint, render_template, request, redirect,
                   url_for, flash, session, jsonify, current_app)
from functools import wraps
from datetime import datetime
import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from flask import send_file

from app import db
from app.models import User, AuditLog, CustomField

auth = Blueprint('auth', __name__)

VALID_ROLES = ('admin', 'support', 'client')


# ── Audit helper ──────────────────────────────────────────────────────────────

def log_action(action, target_type=None, target_id=None, details=None):
    """Write an audit log entry for the currently-logged-in user."""
    try:
        entry = AuditLog(
            user_id=session.get('user_id'),
            username=session.get('username', 'unknown'),
            action=action,
            target_type=target_type,
            target_id=str(target_id) if target_id else None,
            details=details,
            ip_address=request.remote_addr,
        )
        db.session.add(entry)
        db.session.commit()
    except Exception as e:
        current_app.logger.warning(f'[AUDIT] Failed to write log: {e}')


# ── Decorators ────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('auth.login'))
        user = User.query.filter_by(user_id=session['user_id']).first()
        if not user or user.role != 'admin':
            flash('Admin access required.', 'error')
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated


def admin_or_support_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('auth.login'))
        user = User.query.filter_by(user_id=session['user_id']).first()
        if not user or user.role not in ('admin', 'support'):
            return jsonify(success=False, message='Insufficient permissions.'), 403
        return f(*args, **kwargs)
    return decorated


# ── Helper ────────────────────────────────────────────────────────────────────

def get_current_user():
    if 'user_id' in session:
        return User.query.filter_by(user_id=session['user_id']).first()
    return None


# ── Auth routes ───────────────────────────────────────────────────────────────

@auth.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect(url_for('main.dashboard'))

    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')

        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            if not user.is_active:
                error = 'Account is deactivated. Contact an administrator.'
            else:
                session.permanent = True
                session['user_id'] = user.user_id
                session['username'] = user.username
                session['role'] = user.role
                user.last_login = datetime.utcnow()
                db.session.commit()
                log_action('LOGIN', target_type='User', target_id=user.user_id)
                return redirect(url_for('main.dashboard'))
        else:
            error = 'Invalid username or password.'

    return render_template('login.html', error=error)


@auth.route('/logout')
def logout():
    log_action('LOGOUT')
    session.clear()
    return redirect(url_for('auth.login'))


# ── User management ────────────────────────────────────────────────────────────

@auth.route('/users')
@login_required
def users():
    current = get_current_user()
    if not current:
        session.clear()
        return redirect(url_for('auth.login'))
    # Refresh session role from DB in case it's stale
    session['role'] = current.role
    if current.role not in ('admin', 'support'):
        flash('Access denied. Admin or Support role required.', 'error')
        return redirect(url_for('main.dashboard'))
    all_users = User.query.order_by(User.created_at.desc()).all()
    logs = AuditLog.query.order_by(AuditLog.timestamp.desc()).limit(100).all()
    return render_template('users.html', users=all_users, current_user=current, logs=logs)


@auth.route('/users/create', methods=['POST'])
@login_required
def create_user():
    current = get_current_user()
    data     = request.get_json() or request.form
    username = data.get('username', '').strip()
    email    = data.get('email', '').strip()
    password = data.get('password', '')
    role     = data.get('role', 'client')

    if not username or not email or not password:
        return jsonify(success=False, message='Username, email and password are required.'), 400

    # Permission check
    if role == 'support' and not current.can_create_support_user:
        return jsonify(success=False, message='Only admins can create support users.'), 403
    if role == 'client' and not current.can_create_client_user:
        return jsonify(success=False, message='Insufficient permissions to create client users.'), 403
    if role == 'admin' and current.role != 'admin':
        return jsonify(success=False, message='Only admins can create admin users.'), 403
    if role not in VALID_ROLES:
        return jsonify(success=False, message='Invalid role.'), 400

    if User.query.filter_by(username=username).first():
        return jsonify(success=False, message='Username already exists.'), 409
    if User.query.filter_by(email=email).first():
        return jsonify(success=False, message='Email already exists.'), 409

    user = User(username=username, email=email, role=role)
    user.set_password(password)
    db.session.add(user)
    db.session.commit()
    log_action('CREATE_USER', target_type='User', target_id=user.user_id,
               details=f'Created {role} user: {username}')
    return jsonify(success=True, message='User created.', user=user.to_dict())


@auth.route('/users/<int:uid>/update', methods=['POST'])
@login_required
def update_user(uid):
    current = get_current_user()
    if not current or current.role not in ('admin', 'support'):
        return jsonify(success=False, message='Insufficient permissions.'), 403
    user = User.query.get_or_404(uid)
    data = request.get_json() or request.form

    new_username = data.get('username', '').strip()
    new_email    = data.get('email', '').strip()
    new_role     = data.get('role', user.role)
    new_password = data.get('password', '')

    # Permission check for role changes
    if new_role != user.role:
        if new_role == 'support' and not current.can_create_support_user:
            return jsonify(success=False, message='Only admins can assign support role.'), 403
        if new_role == 'admin' and current.role != 'admin':
            return jsonify(success=False, message='Only admins can assign admin role.'), 403
        if new_role not in VALID_ROLES:
            return jsonify(success=False, message='Invalid role.'), 400

    if new_username and new_username != user.username:
        if User.query.filter_by(username=new_username).first():
            return jsonify(success=False, message='Username already taken.'), 409
        user.username = new_username

    if new_email and new_email != user.email:
        if User.query.filter_by(email=new_email).first():
            return jsonify(success=False, message='Email already in use.'), 409
        user.email = new_email

    old_role = user.role
    if new_role in VALID_ROLES:
        user.role = new_role

    if new_password:
        user.set_password(new_password)

    db.session.commit()
    log_action('UPDATE_USER', target_type='User', target_id=user.user_id,
               details=f'Updated user {user.username}; role {old_role}->{user.role}')
    return jsonify(success=True, message='User updated.', user=user.to_dict())


@auth.route('/users/<int:uid>/toggle', methods=['POST'])
@login_required
def toggle_user(uid):
    current = get_current_user()
    if current.role not in ('admin', 'support'):
        return jsonify(success=False, message='Insufficient permissions.'), 403
    user = User.query.get_or_404(uid)
    if user.user_id == session.get('user_id'):
        return jsonify(success=False, message='You cannot deactivate your own account.'), 400
    user.is_active = not user.is_active
    db.session.commit()
    state = 'activated' if user.is_active else 'deactivated'
    log_action(f'{"ACTIVATE" if user.is_active else "DEACTIVATE"}_USER',
               target_type='User', target_id=user.user_id,
               details=f'User {user.username} {state}')
    return jsonify(success=True, message=f'User {state}.', is_active=user.is_active)


@auth.route('/users/<int:uid>/delete', methods=['POST'])
@admin_required
def delete_user(uid):
    user = User.query.get_or_404(uid)
    if user.user_id == session.get('user_id'):
        return jsonify(success=False, message='You cannot delete your own account.'), 400
    uname = user.username
    uid_str = user.user_id
    db.session.delete(user)
    db.session.commit()
    log_action('DELETE_USER', target_type='User', target_id=uid_str,
               details=f'Deleted user {uname}')
    return jsonify(success=True, message='User deleted.')


# ── Excel export ──────────────────────────────────────────────────────────────

@auth.route('/users/export')
@login_required
def export_users():
    current = get_current_user()
    if not current or current.role not in ('admin', 'support'):
        return jsonify(success=False, message='Access denied.'), 403

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Users'

    header_fill = PatternFill('solid', fgColor='1565C0')
    header_font = Font(color='FFFFFF', bold=True)
    headers = ['ID', 'Username', 'Email', 'Role', 'Active', 'Created At', 'Last Login']

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for row, u in enumerate(User.query.order_by(User.id).all(), 2):
        ws.append([
            u.id,
            u.username,
            u.email,
            u.role,
            'Yes' if u.is_active else 'No',
            u.created_at.strftime('%Y-%m-%d %H:%M') if u.created_at else '',
            u.last_login.strftime('%Y-%m-%d %H:%M') if u.last_login else '',
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_action('EXPORT_USERS', target_type='User', details='Exported user list to Excel')
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='users_export.xlsx')


# ── Excel import ──────────────────────────────────────────────────────────────

@auth.route('/users/import', methods=['POST'])
@admin_required
def import_users():
    f = request.files.get('file')
    if not f or not f.filename.endswith('.xlsx'):
        return jsonify(success=False, message='Please upload a valid .xlsx file.'), 400

    wb = openpyxl.load_workbook(f)
    ws = wb.active
    created = skipped = errors = 0
    error_list = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        try:
            username = str(row[0] or '').strip()
            email    = str(row[1] or '').strip()
            password = str(row[2] or '').strip()
            role     = str(row[3] or 'client').strip().lower()
            if role not in VALID_ROLES:
                role = 'client'

            if not username or not email or not password:
                error_list.append(f'Row skipped – missing data: {row}')
                errors += 1
                continue

            if User.query.filter_by(username=username).first() or \
               User.query.filter_by(email=email).first():
                skipped += 1
                continue

            u = User(username=username, email=email, role=role)
            u.set_password(password)
            db.session.add(u)
            created += 1
        except Exception as e:
            error_list.append(str(e))
            errors += 1

    db.session.commit()
    log_action('IMPORT_USERS', details=f'Imported {created} users, {skipped} skipped, {errors} errors')
    return jsonify(success=True,
                   message=f'Import complete: {created} created, {skipped} skipped, {errors} errors.',
                   details=error_list)


# ── Audit log view ─────────────────────────────────────────────────────────────

@auth.route('/api/audit-logs')
@login_required
def get_audit_logs():
    current = get_current_user()
    if not current or current.role not in ('admin', 'support'):
        return jsonify(success=False, message='Access denied.'), 403
    page = request.args.get('page', 1, type=int)
    per_page = 50
    logs = AuditLog.query.order_by(AuditLog.timestamp.desc()).paginate(
        page=page, per_page=per_page, error_out=False)
    return jsonify({
        'logs': [l.to_dict() for l in logs.items],
        'total': logs.total,
        'pages': logs.pages,
        'current_page': page,
    })


# ── Custom Fields ──────────────────────────────────────────────────────────────

@auth.route('/api/custom-fields', methods=['GET'])
@login_required
def list_custom_fields():
    current = get_current_user()
    if not current or not current.can_manage_custom_fields:
        return jsonify(success=False, message='Access denied.'), 403
    fields = CustomField.query.order_by(CustomField.sort_order, CustomField.id).all()
    return jsonify({'fields': [f.to_dict() for f in fields]})


@auth.route('/api/custom-fields', methods=['POST'])
@login_required
def create_custom_field():
    current = get_current_user()
    if not current or not current.can_create_delete_custom_fields:
        return jsonify(success=False, message='Only admins can create custom fields.'), 403

    data = request.get_json() or {}
    name       = data.get('name', '').strip()
    field_type = data.get('field_type', 'text')
    is_required= bool(data.get('is_required', False))
    is_active  = bool(data.get('is_active', True))
    options    = data.get('options', [])
    sort_order = int(data.get('sort_order', 0))

    if not name:
        return jsonify(success=False, message='Field name is required.'), 400
    valid_types = ('text', 'number', 'date', 'single_select', 'multi_select', 'boolean')
    if field_type not in valid_types:
        return jsonify(success=False, message=f'Invalid field type. Must be one of: {", ".join(valid_types)}'), 400
    if CustomField.query.filter_by(name=name).first():
        return jsonify(success=False, message='A field with that name already exists.'), 409

    cf = CustomField(
        name=name,
        field_type=field_type,
        is_required=is_required,
        is_active=is_active,
        options=json.dumps(options) if options else None,
        sort_order=sort_order,
    )
    db.session.add(cf)
    db.session.commit()
    log_action('CREATE_CUSTOM_FIELD', target_type='CustomField', target_id=cf.id,
               details=f'Created field: {name} ({field_type})')
    return jsonify(success=True, field=cf.to_dict())


@auth.route('/api/custom-fields/<int:field_id>', methods=['PUT'])
@login_required
def update_custom_field(field_id):
    current = get_current_user()
    if not current or not current.can_manage_custom_fields:
        return jsonify(success=False, message='Access denied.'), 403

    cf = CustomField.query.get_or_404(field_id)
    data = request.get_json() or {}

    # Support users can toggle active/required but cannot rename or change type
    if current.role == 'support':
        if 'is_active' in data:
            cf.is_active = bool(data['is_active'])
        if 'is_required' in data:
            cf.is_required = bool(data['is_required'])
    else:
        # Admin can change everything
        if 'name' in data:
            new_name = data['name'].strip()
            if new_name != cf.name and CustomField.query.filter_by(name=new_name).first():
                return jsonify(success=False, message='Name already in use.'), 409
            cf.name = new_name
        if 'field_type' in data:
            cf.field_type = data['field_type']
        if 'is_required' in data:
            cf.is_required = bool(data['is_required'])
        if 'is_active' in data:
            cf.is_active = bool(data['is_active'])
        if 'options' in data:
            cf.options = json.dumps(data['options']) if data['options'] else None
        if 'sort_order' in data:
            cf.sort_order = int(data['sort_order'])

    db.session.commit()
    log_action('UPDATE_CUSTOM_FIELD', target_type='CustomField', target_id=cf.id,
               details=f'Updated field: {cf.name}')
    return jsonify(success=True, field=cf.to_dict())


@auth.route('/api/custom-fields/<int:field_id>', methods=['DELETE'])
@login_required
def delete_custom_field(field_id):
    current = get_current_user()
    if not current or not current.can_create_delete_custom_fields:
        return jsonify(success=False, message='Only admins can delete custom fields.'), 403

    cf = CustomField.query.get_or_404(field_id)
    fname = cf.name
    db.session.delete(cf)
    db.session.commit()
    log_action('DELETE_CUSTOM_FIELD', target_type='CustomField', target_id=field_id,
               details=f'Deleted field: {fname}')
    return jsonify(success=True, message=f'Field "{fname}" deleted.')


# ── Profile / change own password ─────────────────────────────────────────────

@auth.route('/profile', methods=['GET', 'POST'])
@login_required
def profile():
    current = get_current_user()
    if not current:
        session.clear()
        return redirect(url_for('auth.login'))
    error = success = None
    if request.method == 'POST':
        old_pw  = request.form.get('old_password', '')
        new_pw  = request.form.get('new_password', '')
        conf_pw = request.form.get('confirm_password', '')

        if not current.check_password(old_pw):
            error = 'Current password is incorrect.'
        elif len(new_pw) < 6:
            error = 'New password must be at least 6 characters.'
        elif new_pw != conf_pw:
            error = 'Passwords do not match.'
        else:
            current.set_password(new_pw)
            db.session.commit()
            log_action('CHANGE_PASSWORD', target_type='User', target_id=current.user_id)
            success = 'Password updated successfully.'

    return render_template('profile.html', current_user=current,
                           error=error, success=success)


# ── Admin: reset another user's password ──────────────────────────────────────

@auth.route('/users/<int:uid>/reset-password', methods=['POST'])
@login_required
def reset_user_password(uid):
    current = get_current_user()
    if not current or current.role not in ('admin', 'support'):
        return jsonify(success=False, message='Insufficient permissions.'), 403
    user = User.query.get_or_404(uid)
    data = request.get_json() or request.form
    new_password = data.get('password', '').strip()
    if len(new_password) < 6:
        return jsonify(success=False, message='Password must be at least 6 characters.'), 400
    # Support users can only reset client passwords
    if current.role == 'support' and user.role != 'client':
        return jsonify(success=False, message='Support users can only reset client passwords.'), 403
    user.set_password(new_password)
    db.session.commit()
    log_action('RESET_PASSWORD', target_type='User', target_id=user.user_id,
               details=f'Password reset for user {user.username}')
    return jsonify(success=True, message=f'Password reset for {user.username}.')


# ── Role Permissions API ───────────────────────────────────────────────────────

@auth.route('/api/permissions', methods=['GET'])
@login_required
def get_permissions():
    """Return the full permission matrix (role → permission → granted)."""
    from app.models import RolePermission
    rows = RolePermission.query.all()
    data = {}
    for row in rows:
        data.setdefault(row.role, {})[row.permission] = row.granted
    return jsonify({'permissions': data})


@auth.route('/api/permissions', methods=['POST'])
@login_required
def save_permissions():
    """Save permission overrides.
    Expects JSON: { changes: [ {role, permission, granted}, ... ] }
    Only admins may save.
    """
    from app.models import RolePermission

    current = get_current_user()
    if not current or current.role != 'admin':
        return jsonify(success=False, message='Admin access required.'), 403

    data    = request.get_json() or {}
    changes = data.get('changes', [])

    # Hard-locked permissions that can never be changed via the UI
    LOCKED = {
        ('admin',   'can_create_admin_user'),
        ('support', 'can_create_admin_user'),
        ('client',  'can_create_admin_user'),
        ('admin',   'can_delete_users'),
        ('support', 'can_delete_users'),
        ('client',  'can_delete_users'),
        ('admin',   'can_create_delete_custom_fields'),  # admin always True
        ('support', 'can_create_delete_custom_fields'),  # always False for support
        ('client',  'can_create_delete_custom_fields'),  # always False for client
    }

    saved = 0
    for ch in changes:
        role       = ch.get('role', '')
        permission = ch.get('permission', '')
        granted    = bool(ch.get('granted', False))

        if (role, permission) in LOCKED:
            continue  # silently skip locked cells
        if role not in ('admin', 'support', 'client'):
            continue

        row = RolePermission.query.filter_by(role=role, permission=permission).first()
        if row:
            row.granted = granted
        else:
            row = RolePermission(role=role, permission=permission, granted=granted)
            db.session.add(row)
        saved += 1

    try:
        db.session.commit()
        log_action('SAVE_PERMISSIONS', details=f'Updated {saved} permission override(s)')
        return jsonify(success=True, saved=saved)
    except Exception as e:
        db.session.rollback()
        return jsonify(success=False, message=str(e)), 500
